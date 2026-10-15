import asyncio
import json
import logging
import time
from collections.abc import AsyncIterator
from io import BytesIO
from uuid import UUID, uuid4

from fastapi import APIRouter, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from langgraph.types import Command

from app.graph.state import Plan
from app.graph.workflow import blog_writer
from app.services.rag_service import (
    delete_document_batch,
    ingest_documents,
    retrieve_context,
)
from .schemas import OutlineResponse, OutlineRevisionRequest

router = APIRouter()
logger = logging.getLogger(__name__)

MAX_FILE_SIZE = 10 * 1024 * 1024
MAX_TOTAL_CHARS = 200_000

# Checkpoints store workflow state; this map stores the outline metadata used by the API.
pending_plans: dict[str, Plan] = {}
pending_document_batches: dict[str, str] = {}
active_tasks: dict[str, asyncio.Task] = {}


def _delete_document_batch(batch_id: str) -> None:
    try:
        delete_document_batch(batch_id)
    except Exception:
        logger.exception("Could not delete Chroma document batch %s", batch_id)


def _event(**payload: object) -> str:
    return json.dumps(payload, ensure_ascii=False) + "\n"


def _extract_file_text(filename: str, content: bytes) -> str:
    if filename.endswith((".txt", ".md", ".csv")):
        return content.decode("utf-8")

    if filename.endswith(".pdf"):
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(content))
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    if filename.endswith(".docx"):
        from docx import Document

        doc = Document(BytesIO(content))
        paragraphs = [paragraph.text for paragraph in doc.paragraphs]
        tables = [
            "\n".join("\t".join(cell.text for cell in row.cells) for row in table.rows)
            for table in doc.tables
        ]
        return "\n".join([*paragraphs, *tables])

    raise HTTPException(
        status_code=415,
        detail=f"Unsupported file format: {filename}",
    )


async def extract_documents(files: list[UploadFile] | None) -> str:
    if not files:
        return ""

    documents: list[str] = []
    total_chars = 0

    for file in files:
        filename = (file.filename or "").lower()
        try:
            content = await file.read(MAX_FILE_SIZE + 1)
            if len(content) > MAX_FILE_SIZE:
                raise HTTPException(
                    status_code=413,
                    detail=f"{filename} exceeds the 10 MB limit.",
                )

            text = _extract_file_text(filename, content).strip()
            if not text:
                raise HTTPException(
                    status_code=422,
                    detail=f"No extractable text was found in {filename}.",
                )

            remaining_chars = MAX_TOTAL_CHARS - total_chars
            if remaining_chars > 0:
                document = f"\n--- Document: {filename} ---\n{text}"
                document = document[:remaining_chars]
                documents.append(document)
                total_chars += len(document)
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(
                status_code=422,
                detail=f"Could not read {filename}: {exc}",
            ) from exc
        finally:
            await file.close()

    return "\n".join(documents)


async def run_graph(approval_id: str) -> AsyncIterator[str]:
    current_task = asyncio.current_task()
    if current_task is not None:
        active_tasks[approval_id] = current_task
    started_at = time.perf_counter()
    plan = pending_plans.get(approval_id)
    plan_titles = {
        task.id: task.title for task in (plan.tasks if plan is not None else [])
    }
    total_sections = len(plan_titles)
    completed_sections = 0
    config = {"configurable": {"thread_id": approval_id}}

    def emit(
        name: str,
        message: str,
        *,
        section: str | None = None,
        progress: int = 0,
        event_id: str | None = None,
    ) -> str:
        return _event(
            id=event_id or name,
            name=name,
            message=message,
            section=section,
            completed=True,
            progress=progress,
            elapsed_ms=round((time.perf_counter() - started_at) * 1000),
        )

    yield emit("start", "Resuming approved blog generation")

    try:
        async for update in blog_writer.astream(
            Command(resume={"action": "approve"}),
            config=config,
            stream_mode="updates",
        ):
            for name, result in update.items():
                if not isinstance(result, dict):
                    continue

                if name == "reducer":
                    final_blog = result.get("final", "")
                    yield _event(
                        finished=True,
                        blog=final_blog,
                        blog_title=(
                            final_blog.split("\n", 1)[0]
                            .removeprefix("# ")
                            .strip()
                        ),
                        elapsed_ms=round((time.perf_counter() - started_at) * 1000),
                    )
                    continue

                if name == "worker":
                    sections = result.get("sections", [])
                    section_id = sections[0][0] if sections else 0
                    completed_sections += 1
                    section_title = plan_titles.get(
                        section_id, f"Section {section_id}"
                    )
                    progress = 35 + round(
                        completed_sections / max(total_sections, 1) * 50
                    )
                    yield emit(
                        name,
                        f"Finished section {completed_sections} of {total_sections}",
                        section=section_title,
                        progress=progress,
                        event_id=f"section-{section_id}",
                    )
                else:
                    yield emit(name, f"Completed {name.replace('_', ' ')}")
    except Exception as exc:
        yield _event(
            error="Blog generation failed",
            detail=str(exc),
            elapsed_ms=round((time.perf_counter() - started_at) * 1000),
        )
    finally:
        if active_tasks.get(approval_id) is current_task:
            active_tasks.pop(approval_id, None)
        pending_plans.pop(approval_id, None)
        batch_id = pending_document_batches.pop(approval_id, None)
        if batch_id:
            _delete_document_batch(batch_id)


def _outline_response(approval_id: str, topic: str, plan: Plan) -> OutlineResponse:
    return OutlineResponse(
        approval_id=approval_id,
        topic=topic,
        blog_title=plan.blog_title,
        sections=[task.model_dump() for task in plan.tasks],
    )


@router.get("/")
def health() -> dict[str, str]:
    return {"content": "server running"}


@router.post("/generate", response_model=OutlineResponse)
async def generate_blog(
    request_id: str = Form(..., min_length=32, max_length=36),
    topic: str = Form(..., min_length=3, max_length=500),
    supporting_documents: list[UploadFile] | None = File(default=None),
) -> OutlineResponse:
    topic = topic.strip()
    try:
        approval_id = UUID(request_id).hex
    except ValueError as exc:
        raise HTTPException(status_code=422, detail="Invalid request ID.") from exc
    batch_id = ""
    current_task = asyncio.current_task()
    if current_task is not None:
        active_tasks[approval_id] = current_task
    initial_state = {
        "topic": topic,
        "supporting_context": "",
        "document_batch_id": "",
        "document_summary": "",
        "plan": None,
        "mode": "",
        "feedback_action": "",
        "revision_instructions": "",
        "needs_research": False,
        "is_document_relevant": False,
        "queries": [],
        "evidence": [],
        "sections": [],
        "final": "",
    }

    if supporting_documents:
        try:
            extracted_text = await extract_documents(supporting_documents)
        except asyncio.CancelledError:
            if active_tasks.get(approval_id) is current_task:
                active_tasks.pop(approval_id, None)
            raise
        finally:
            if active_tasks.get(approval_id) is current_task:
                active_tasks.pop(approval_id, None)
        if extracted_text.strip():
            batch_id = uuid4().hex
            try:
                ingest_documents(extracted_text, batch_id)
                supporting_context = retrieve_context(topic, batch_id)
            except Exception as exc:
                _delete_document_batch(batch_id)
                raise HTTPException(
                    status_code=422,
                    detail=f"Could not process the supporting documents: {exc}",
                ) from exc
            initial_state["supporting_context"] = supporting_context or extracted_text
            initial_state["document_batch_id"] = batch_id

    config = {"configurable": {"thread_id": approval_id}}
    if current_task is not None:
        active_tasks[approval_id] = current_task
    try:
        result = await blog_writer.ainvoke(initial_state, config)
    except asyncio.CancelledError:
        if batch_id:
            _delete_document_batch(batch_id)
        raise
    except Exception as exc:
        if batch_id:
            _delete_document_batch(batch_id)
        raise HTTPException(
            status_code=500,
            detail=f"Could not generate the outline: {exc}",
        ) from exc
    finally:
        if active_tasks.get(approval_id) is current_task:
            active_tasks.pop(approval_id, None)

    plan_data = result.get("plan")
    if plan_data is None or not result.get("__interrupt__"):
        if batch_id:
            _delete_document_batch(batch_id)
        raise HTTPException(
            status_code=500,
            detail="The workflow did not pause for outline approval.",
        )

    plan = Plan.model_validate(plan_data)
    pending_plans[approval_id] = plan
    if batch_id:
        pending_document_batches[approval_id] = batch_id
    return _outline_response(approval_id, topic, plan)


@router.get("/generate")
async def approve_blog(approval_id: str = Query(..., min_length=8)) -> StreamingResponse:
    if approval_id not in pending_plans:
        return StreamingResponse(
            iter([_event(error="This approval has expired or was already used.")]),
            media_type="application/x-ndjson",
            status_code=410,
        )

    return StreamingResponse(
        run_graph(approval_id),
        media_type="application/x-ndjson",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


@router.post("/generate/{approval_id}/revise", response_model=OutlineResponse)
async def revise_outline(
    approval_id: str,
    request: OutlineRevisionRequest,
) -> OutlineResponse:
    if approval_id not in pending_plans:
        raise HTTPException(
            status_code=410,
            detail="This outline has expired or was already used.",
        )

    config = {"configurable": {"thread_id": approval_id}}
    current_task = asyncio.current_task()
    if current_task is not None:
        active_tasks[approval_id] = current_task
    try:
        result = await blog_writer.ainvoke(
            Command(
                resume={
                    "action": "revise",
                    "instructions": request.instructions.strip(),
                }
            ),
            config=config,
        )
    except asyncio.CancelledError:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=500,
            detail=f"Could not regenerate the outline: {exc}",
        ) from exc
    finally:
        if active_tasks.get(approval_id) is current_task:
            active_tasks.pop(approval_id, None)

    plan_data = result.get("plan")
    if plan_data is None or not result.get("__interrupt__"):
        raise HTTPException(
            status_code=500,
            detail="The revised outline did not pause for approval.",
        )

    plan = Plan.model_validate(plan_data)
    pending_plans[approval_id] = plan
    return _outline_response(approval_id, result.get("topic", ""), plan)


@router.post("/generate/{approval_id}/cancel")
async def cancel_generation(approval_id: str) -> dict[str, str]:
    active_task = active_tasks.get(approval_id)
    if approval_id not in pending_plans and active_task is None:
        raise HTTPException(
            status_code=410,
            detail="This outline has expired or was already used.",
        )

    # Approval IDs identify one LangGraph thread and its document batch.
    if active_task is not None and active_task is not asyncio.current_task():
        active_task.cancel()
        await asyncio.gather(active_task, return_exceptions=True)

    pending_plans.pop(approval_id, None)
    batch_id = pending_document_batches.pop(approval_id, None)
    if batch_id:
        _delete_document_batch(batch_id)
    return {"status": "cancelled"}
